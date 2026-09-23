#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import sys, os, argparse, io

if sys.platform == "win32":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")


def sanitize_query(q):
    if not q:
        return q
    # Replace safe word aliases for PowerShell safety
    replacements = {
        ' gt ': ' > ',
        ' gte ': ' >= ',
        ' lt ': ' < ',
        ' lte ': ' <= ',
        ' eq ': ' == ',
        ' ne ': ' != '
    }
    res = q
    for k, v in replacements.items():
        res = res.replace(k, v)
    return res

def inspect_excel(path, args):
    import pandas as pd
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        sheets = {"CSV": path}
    else:
        xls = pd.ExcelFile(path)
        sheets = {name: name for name in xls.sheet_names}

    if args.summary or (not args.sheet and not args.search and not args.markdown):
        print(f"# Summary: {os.path.basename(path)}")
        print(f"- Format: {ext.upper().replace('.', '')}")
        sheet_names = ", ".join([f"`{s}`" for s in list(sheets.keys())[:10]])
        print(f"- Sheets ({len(sheets)}): {sheet_names}")
        print("\n---")
        for sname in list(sheets.keys())[:3]:
            try:
                if ext == ".csv":
                    df = pd.read_csv(path, nrows=args.head or 5)
                    total_rows = sum(1 for _ in open(path, "rb")) - 1
                else:
                    df = pd.read_excel(path, sheet_name=sname, nrows=args.head or 5)
                    total_rows = len(pd.read_excel(path, sheet_name=sname, usecols=[0]))
                cols_preview = ", ".join([f"`{c}` ({str(df[c].dtype)})" for c in df.columns[:15]])
                print(f"### Sheet: `{sname}` ({total_rows:,} rows x {len(df.columns)} cols)")
                print(f"**Columns:** {cols_preview}")
                print("\n**Sample Preview:**")
                print(df.head(args.head or 3).to_markdown(index=False))
                print()
            except Exception as e:
                print(f"- Could not preview `{sname}`: {e}")
        return

    if args.sheet:
        try:
            df = pd.read_csv(path) if ext == ".csv" else pd.read_excel(path, sheet_name=args.sheet)
            if args.query:
                clean_q = sanitize_query(args.query)
                filtered_df = df.query(clean_q)
                print(f"### Query: `{args.query}` on `{args.sheet}`")
                print(f"- Matches: {len(filtered_df):,} / {len(df):,} rows")
                print(filtered_df.head(args.head or 10).to_markdown(index=False))
            else:
                print(f"### Sheet: `{args.sheet}` ({len(df):,} rows x {len(df.columns)} cols)")
                print(df.head(args.head or 10).to_markdown(index=False))
        except Exception as e:
            print(f"Error reading sheet `{args.sheet}`: {e}")
        return

    if args.search:
        print(f"# Search Results for: `{args.search}` in `{os.path.basename(path)}`")
        found = 0
        for sname in sheets:
            try:
                df = pd.read_csv(path) if ext == ".csv" else pd.read_excel(path, sheet_name=sname)
                mask = df.astype(str).apply(lambda col: col.str.contains(args.search, case=False, na=False)).any(axis=1)
                matches = df[mask]
                if not matches.empty:
                    print(f"### Match in `{sname}`: {len(matches)} rows")
                    print(matches.head(5).to_markdown(index=False))
                    found += len(matches)
            except Exception:
                continue
        if found == 0:
            print("No matches found.")
        return

def inspect_pdf(path, args):
    from pypdf import PdfReader
    reader = PdfReader(path)
    total_pages = len(reader.pages)
    if args.summary or (not args.pages and not args.search and not args.markdown):
        meta = reader.metadata or {}
        print(f"# PDF Summary: {os.path.basename(path)}")
        print(f"- Total Pages: {total_pages}")
        if meta.title:
            print(f"- Title: {meta.title}")
        if meta.author:
            print(f"- Author: {meta.author}")
        try:
            outline = reader.outline
            if outline:
                print("\n### Table of Contents / Outline:")
                def print_outline(items, level=0):
                    for item in items[:25]:
                        if isinstance(item, list):
                            print_outline(item, level + 1)
                        else:
                            title = getattr(item, "title", str(item))
                            print(f"{'  ' * level}- {title}")
                print_outline(outline)
        except Exception:
            pass
        print("\n### Sample Preview (Page 1 excerpt):")
        if total_pages > 0:
            text = reader.pages[0].extract_text() or ""
            preview = text.strip()[:400].replace("\n", " ")
            print(f"> {preview} ...")
        return

    if args.pages:
        if "-" in args.pages:
            p_start, p_end = map(int, args.pages.split("-"))
            pages_to_extract = list(range(p_start, min(p_end + 1, total_pages + 1)))
        else:
            pages_to_extract = [int(args.pages)]
        print(f"# PDF Extract: {os.path.basename(path)} (Pages: {args.pages})")
        for p_num in pages_to_extract:
            idx = p_num - 1
            if 0 <= idx < total_pages:
                text = reader.pages[idx].extract_text() or ""
                print(f"\n## Page {p_num}\n")
                print(text.strip())
        return

    if args.search:
        print(f"# Search Results for: `{args.search}` in `{os.path.basename(path)}`")
        found = 0
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if args.search.lower() in text.lower():
                found += 1
                lines = [ln.strip() for ln in text.splitlines() if args.search.lower() in ln.lower()]
                first_match = lines[0] if lines else "Found"
                print(f"- Page {i + 1}: {first_match}")
                if found >= 15:
                    print("... and more matches.")
                    break
        if found == 0:
            print("No matches found.")
        return

def inspect_docx(path, args):
    import docx
    doc = docx.Document(path)
    paragraphs = doc.paragraphs
    tables = doc.tables
    total_words = sum(len(p.text.split()) for p in paragraphs)
    if args.summary or (not args.search and not args.markdown):
        print(f"# Word Document Summary: {os.path.basename(path)}")
        print(f"- Paragraphs: {len(paragraphs):,}")
        print(f"- Approx Words: {total_words:,}")
        print(f"- Tables: {len(tables)}")
        headings = [p.text.strip() for p in paragraphs if p.style.name.startswith("Heading") and p.text.strip()]
        if headings:
            print("\n### Outline (Headings):")
            for h in headings[:25]:
                print(f"- {h}")
        else:
            print("\n### Preview (First 3 paragraphs):")
            for p in [p for p in paragraphs if p.text.strip()][:3]:
                print(f"> {p.text.strip()}")
        return

    if args.search:
        print(f"# Search Results for: `{args.search}` in `{os.path.basename(path)}`")
        found = 0
        for idx, p in enumerate(paragraphs):
            if args.search.lower() in p.text.lower():
                found += 1
                print(f"- Paragraph {idx + 1}: {p.text.strip()}")
                if found >= 15:
                    break
        if found == 0:
            print("No matches found.")
        return

def inspect_pptx(path, args):
    import pptx
    prs = pptx.Presentation(path)
    total_slides = len(prs.slides)
    if args.summary or (not args.search and not args.markdown):
        print(f"# PowerPoint Summary: {os.path.basename(path)}")
        print(f"- Total Slides: {total_slides}")
        print("\n### Slide Outline:")
        for idx, slide in enumerate(prs.slides):
            title = "Untitled Slide"
            if slide.shapes.title and slide.shapes.title.text:
                title = slide.shapes.title.text.strip().replace("\n", " ")
            print(f"- Slide {idx + 1}: {title}")
        return

    if args.search:
        print(f"# Search Results for: `{args.search}` in `{os.path.basename(path)}`")
        found = 0
        for idx, slide in enumerate(prs.slides):
            texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
            full_text = " ".join(texts)
            if args.search.lower() in full_text.lower():
                found += 1
                title = slide.shapes.title.text.strip() if slide.shapes.title else f"Slide {idx + 1}"
                print(f"- Slide {idx + 1} ({title}): Found match")
        if found == 0:
            print("No matches found.")
        return

def convert_markdown(path):
    import markitdown
    md = markitdown.MarkItDown()
    result = md.convert(path)
    print(result.text_content)

def main():
    parser = argparse.ArgumentParser(description="Surgical Document & Office Token Optimizer")
    parser.add_argument("path", help="Path to document")
    parser.add_argument("--summary", action="store_true", help="Print summary")
    parser.add_argument("--sheet", help="Sheet name")
    parser.add_argument("--head", type=int, default=5, help="Number of rows")
    parser.add_argument("--query", help="Pandas query")
    parser.add_argument("--pages", help="Page range (e.g. 1-3)")
    parser.add_argument("--search", help="Search keyword")
    parser.add_argument("--markdown", action="store_true", help="Convert to markdown")
    args = parser.parse_args()

    if not os.path.exists(args.path):
        print(f"Error: File not found: {args.path}", file=sys.stderr)
        sys.exit(1)

    ext = os.path.splitext(args.path)[1].lower()
    if args.markdown:
        convert_markdown(args.path)
        return

    if ext in [".xlsx", ".xls", ".csv"]:
        inspect_excel(args.path, args)
    elif ext == ".pdf":
        inspect_pdf(args.path, args)
    elif ext == ".docx":
        inspect_docx(args.path, args)
    elif ext == ".pptx":
        inspect_pptx(args.path, args)
    else:
        convert_markdown(args.path)

if __name__ == "__main__":
    main()
